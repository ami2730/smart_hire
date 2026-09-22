"""
Tests for document extraction — Phase 2.

Verifies:
- PDF text extraction via PyMuPDF
- DOCX text extraction via python-docx
- Extension validation
- Header / magic-byte verification
- File size restrictions
- REST API: POST /api/v1/resume/extract
"""

from __future__ import annotations

import io
import pytest
import docx
import pymupdf as fitz
from fastapi.testclient import TestClient

from app.core.exceptions import (
    FileTooLargeError,
    UnsupportedFileTypeError,
)
from app.main import app
from app.services.document_service import DocumentService


@pytest.fixture(scope="module")
def client() -> TestClient:
    return TestClient(app)


@pytest.fixture
def document_service() -> DocumentService:
    return DocumentService()


def create_sample_pdf_bytes(text: str = "Jane Doe\nBackend Developer with Python and PostgreSQL.") -> bytes:
    """Create a minimal in-memory PDF using PyMuPDF."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), text)
    pdf_bytes = doc.write()
    doc.close()
    return pdf_bytes


def create_sample_docx_bytes(text: str = "John Smith\nFull-Stack Engineer with React and Node.js.") -> bytes:
    """Create a minimal in-memory DOCX using python-docx."""
    doc = docx.Document()
    doc.add_heading("Resume", level=1)
    doc.add_paragraph(text)
    table = doc.add_table(rows=1, cols=2)
    row = table.rows[0]
    row.cells[0].text = "Skill"
    row.cells[1].text = "FastAPI"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestDocumentServiceUnit:
    """Unit tests for DocumentService."""

    def test_extract_pdf_success(self, document_service: DocumentService) -> None:
        pdf_bytes = create_sample_pdf_bytes("Alice Engineer\nPython and Docker specialist.")
        result = document_service.extract_from_bytes(
            content=pdf_bytes,
            filename="resume.pdf",
            content_type="application/pdf",
        )
        assert result.document_type == "pdf"
        assert result.page_count == 1
        assert "Alice Engineer" in result.text
        assert "Python and Docker" in result.text

    def test_extract_docx_success(self, document_service: DocumentService) -> None:
        docx_bytes = create_sample_docx_bytes("Bob Developer\nExpert in TypeScript and Next.js.")
        result = document_service.extract_from_bytes(
            content=docx_bytes,
            filename="resume.docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        assert result.document_type == "docx"
        assert result.page_count is None
        assert "Bob Developer" in result.text
        assert "FastAPI" in result.text  # From table

    def test_unsupported_extension_raises_error(self, document_service: DocumentService) -> None:
        with pytest.raises(UnsupportedFileTypeError) as exc_info:
            document_service.extract_from_bytes(
                content=b"some text",
                filename="resume.txt",
            )
        assert "Unsupported file extension" in str(exc_info.value.message)

    def test_fake_pdf_header_raises_error(self, document_service: DocumentService) -> None:
        fake_pdf = b"This is not really a PDF file."
        with pytest.raises(UnsupportedFileTypeError) as exc_info:
            document_service.extract_from_bytes(
                content=fake_pdf,
                filename="resume.pdf",
            )
        assert "does not match PDF specification" in str(exc_info.value.message)

    def test_fake_docx_header_raises_error(self, document_service: DocumentService) -> None:
        fake_docx = b"This is not a zip docx."
        with pytest.raises(UnsupportedFileTypeError) as exc_info:
            document_service.extract_from_bytes(
                content=fake_docx,
                filename="resume.docx",
            )
        assert "does not match Microsoft OpenXML specification" in str(exc_info.value.message)

    def test_oversized_file_raises_error(self, document_service: DocumentService) -> None:
        # 11MB file exceeding 10MB limit
        oversized = b"%PDF-" + (b"0" * (11 * 1024 * 1024))
        with pytest.raises(FileTooLargeError) as exc_info:
            document_service.extract_from_bytes(
                content=oversized,
                filename="large_resume.pdf",
            )
        assert "exceeds maximum allowed size" in str(exc_info.value.message)


class TestResumeExtractEndpoint:
    """Integration tests for POST /api/v1/resume/extract."""

    def test_extract_valid_pdf_endpoint(self, client: TestClient) -> None:
        pdf_bytes = create_sample_pdf_bytes("Backend Developer with 4 years Python and PostgreSQL.")
        files = {"file": ("my_resume.pdf", pdf_bytes, "application/pdf")}
        response = client.post("/api/v1/resume/extract", files=files)

        assert response.status_code == 200
        body = response.json()
        assert body["success"] is True
        assert body["document_type"] == "pdf"
        assert body["page_count"] == 1
        assert "Python and PostgreSQL" in body["text"]

    def test_extract_valid_docx_endpoint(self, client: TestClient) -> None:
        docx_bytes = create_sample_docx_bytes("Software Engineer specializing in React and FastAPI.")
        files = {
            "file": (
                "resume.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        response = client.post("/api/v1/resume/extract", files=files)

        assert response.status_code == 200
        body = response.json()
        assert body["success"] is True
        assert body["document_type"] == "docx"
        assert body["page_count"] is None
        assert "Software Engineer" in body["text"]

    def test_extract_unsupported_type_returns_422(self, client: TestClient) -> None:
        files = {"file": ("malicious.exe", b"binary content", "application/x-msdownload")}
        response = client.post("/api/v1/resume/extract", files=files)

        assert response.status_code == 422
        body = response.json()
        assert body["success"] is False
        assert body["error"]["code"] == "UNSUPPORTED_FILE_TYPE"

    def test_extract_corrupted_pdf_returns_422(self, client: TestClient) -> None:
        files = {"file": ("corrupt.pdf", b"NOT_A_REAL_PDF", "application/pdf")}
        response = client.post("/api/v1/resume/extract", files=files)

        assert response.status_code == 422
        body = response.json()
        assert body["success"] is False
        assert body["error"]["code"] == "UNSUPPORTED_FILE_TYPE"
