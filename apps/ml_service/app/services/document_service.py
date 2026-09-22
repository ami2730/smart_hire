"""
Document Processing Service.

Handles text extraction from PDF (via PyMuPDF) and DOCX (via python-docx) documents.
Performs validation, format detection, content extraction, and content normalization.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Literal

import docx
import pymupdf as fitz

from app.core.exceptions import DocumentExtractionError
from app.core.logging import get_logger
from app.utils.text import clean_whitespace
from app.utils.validation import validate_document_type, validate_file_size

logger = get_logger(__name__)


@dataclass(frozen=True)
class DocumentExtractionResult:
    """Represents the output of document text extraction."""

    document_type: Literal["pdf", "docx"]
    text: str
    page_count: int | None


class DocumentService:
    """Service for extracting text from uploaded resume documents."""

    def extract_from_bytes(
        self,
        content: bytes,
        filename: str,
        content_type: str | None = None,
    ) -> DocumentExtractionResult:
        """Validate uploaded file and extract text and page metadata.

        Args:
            content: Raw binary content of the file.
            filename: Original filename with extension.
            content_type: Optional HTTP Content-Type header.

        Returns:
            DocumentExtractionResult with document type, normalized text, and page count.

        Raises:
            UnsupportedFileTypeError: If file type or signature is unsupported.
            FileTooLargeError: If file exceeds configured size limit.
            DocumentExtractionError: If document is corrupt or cannot be parsed.
        """
        # 1. Size validation
        validate_file_size(content)

        # 2. Format validation (extension, magic bytes, MIME)
        doc_type = validate_document_type(filename, content, content_type)

        # 3. Extraction based on verified format
        if doc_type == "pdf":
            return self._extract_pdf(content)
        elif doc_type == "docx":
            return self._extract_docx(content)

        raise DocumentExtractionError(f"Unsupported document format: {doc_type}")

    def _extract_pdf(self, content: bytes) -> DocumentExtractionResult:
        """Extract text and page count from PDF bytes using PyMuPDF."""
        try:
            doc = fitz.open(stream=content, filetype="pdf")
            page_count = doc.page_count
            pages_text: list[str] = []

            for page in doc:
                text = page.get_text()
                if text and text.strip():
                    pages_text.append(text.strip())

            doc.close()
            combined_text = "\n\n".join(pages_text)
            normalized_text = clean_whitespace(combined_text)

            logger.info(
                "PDF extracted successfully | pages=%d | characters=%d",
                page_count,
                len(normalized_text),
            )
            return DocumentExtractionResult(
                document_type="pdf",
                text=normalized_text,
                page_count=page_count,
            )
        except Exception as exc:
            logger.error("Failed to extract text from PDF: %s", exc)
            raise DocumentExtractionError(
                f"Failed to read PDF document: {exc}"
            ) from exc

    def _extract_docx(self, content: bytes) -> DocumentExtractionResult:
        """Extract text from DOCX bytes using python-docx."""
        try:
            file_stream = io.BytesIO(content)
            doc = docx.Document(file_stream)
            text_parts: list[str] = []

            # Extract paragraph text
            for para in doc.paragraphs:
                cleaned = para.text.strip()
                if cleaned:
                    text_parts.append(cleaned)

            # Extract text from tables (common in resume layouts)
            for table in doc.tables:
                for row in table.rows:
                    row_parts = [
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    ]
                    if row_parts:
                        text_parts.append(" | ".join(row_parts))

            combined_text = "\n\n".join(text_parts)
            normalized_text = clean_whitespace(combined_text)

            logger.info(
                "DOCX extracted successfully | characters=%d",
                len(normalized_text),
            )
            return DocumentExtractionResult(
                document_type="docx",
                text=normalized_text,
                page_count=None,  # Page count cannot be reliably determined for DOCX without Word rendering engine
            )
        except Exception as exc:
            logger.error("Failed to extract text from DOCX: %s", exc)
            raise DocumentExtractionError(
                f"Failed to read DOCX document: {exc}"
            ) from exc
